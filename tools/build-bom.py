"""Builds docs/Villa-Somalia-SCADA-BOM.xlsx – project bill of materials for the
Schneider Electric architecture on drawings BFE/PR/26/01-SCADA-001 and -002.

Usage: python3 tools/build-bom.py   (then recalculate with LibreOffice)
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "docs" / "Villa-Somalia-SCADA-BOM.xlsx"

# Unit prices copied from the Schneider offer workbook "Bill of Material for Villa Somalia",
# sheet "BOM by Reference" (unit price column). Currency as in the offer (to be confirmed).
OFFER_PRICE = {
    "HMIP67GD0F88BE7N00": 6750, "HMIEMSE3PRT4KA": 4699.913954249348, "HMIEMSEBT4KA": 2749.032547941522,
    "BMEXBP0800": 227.682, "BMXCPS3020": 320.0265, "BMEP582040": 2337.7635, "BMENOC0311": 887.94,
    "BMENOC0321": 1014.795, "BMXDDI3202K": 212.1705, "BMXFCW303": 51.5835, "BMXDDO3202K": 250.218,
    "BMEXBP1200": 322.7805, "BMECRA31210": 789.4575, "BMXNOM0200": 414.8775, "VW3A8306R10": 7.965,
    "MCSESM083F23F1": 1609.695, "NTSNEC1200K": 308.772, "NTSPFB1002HK": 88.695, "NTSXTB02030H": 18,
    "NTSDDI1602XK": 122.2884, "NTSXTB18001XH": 63, "NTSDDO1602K": 137.268, "NTSXTB18001H": 58,
    "NTSFMB0120K": 800, "NTSXTB12001H": 40, "PAS600": 1106, "130729": 99, "EMS59440": 230,
    "SVC-DESIGN-COMM": 37567,
}
# Offer "Total Quantity" per reference (sheet "BOM by Reference"; REL50458, METSEPM5320 and
# PVSCL60E from the offer PDF bill of material, which the workbook omits).
OFFER_QTY = {
    "HMIP67GD0F88BE7N00": 2, "HMIEMSE3PRT4KA": 2, "HMIEMSEBT4KA": 1, "BMEXBP0800": 2, "BMXCPS3020": 3,
    "BMEP582040": 2, "BMENOC0311": 2, "BMENOC0321": 1, "BMXDDI3202K": 5, "BMXFCW303": 7, "BMXDDO3202K": 4,
    "BMEXBP1200": 1, "BMECRA31210": 1, "BMXNOM0200": 1, "VW3A8306R10": 1, "MCSESM083F23F1": 2,
    "NTSNEC1200K": 7, "NTSPFB1002HK": 7, "NTSXTB02030H": 12, "NTSDDI1602XK": 19, "NTSXTB18001XH": 18,
    "NTSDDO1602K": 7, "NTSXTB18001H": 6, "NTSFMB0120K": 6, "NTSXTB12001H": 6, "PAS600": 7, "130729": 7,
    "EMS59440": 9, "REL50458": 8, "METSEPM5320": 5, "PVSCL60E": 4, "SVC-DESIGN-COMM": 1,
}

OFF, ADD, TBC, BYO = "Offered", "Added", "Added – ref. TBC", "By others"

# (loc code, location name, [ (tag, reference, description, qty, unit, source, status, remark) ])
LOCATIONS = [
 ("A", "Control room – SCADA, servers and network (Level 2/3)", [
  ("HMI-01/02", "HMIP67GD0F88BE7N00", "Harmony P6 industrial PC – operator stations (configuration per offer)", 2, "pc", "Offer", OFF, "Offer gives no description; confirm screen size, CPU and OS."),
  ("HMI-01/02", "HMIEMSE3PRT4KA", "EcoStruxure Machine SCADA Expert 3rd-party runtime 4K tags, digital licence", 2, "pc", "Offer", OFF, "4K tags per station; I/O list holds 958 signals plus internal tags."),
  ("ENG-01", "HMIEMSEBT4KA", "EcoStruxure Machine SCADA Expert build-time 4K tags, digital licence", 1, "pc", "Offer", OFF, ""),
  ("SRV-01", "TBC-SRV-01", "SCADA server / historian, rack industrial server, RAID 1, redundant PSU, 13-month storage", 1, "pc", "Sheet 001", TBC, "Not in offer. Needed for historian and alarm/event archive."),
  ("ENG-01", "TBC-ENG-01", "Engineering workstation (laptop) with EcoStruxure Control Expert licence for M580", 1, "pc", "Sheet 001", TBC, "Not in offer. Needed to program and maintain the M580."),
  ("GPS-01", "TBC-GPS-01", "GPS master clock with antenna, NTP / PTP server", 1, "pc", "Sheet 001", TBC, "Time source for 1 ms relay event stamping."),
  ("FW-01", "TBC-FW-01", "Industrial firewall, DMZ between Level 2 and Level 3", 1, "pc", "Sheet 001", TBC, "Only conduit to remote access (client policy TBC)."),
  ("SW1", "MCSESM083F23F1", "Modicon Extended Managed Switch 8TX – SCADA LAN (SW1)", 1, "pc", "Offer", OFF, "Offer has 2; the second is replaced by an FO variant for SW2."),
  ("UPS-01", "SRT6KRMXLI", "APC Smart-UPS SRT 6 kVA rack/tower, 230 V – SCADA, PLC and switches", 1, "pc", "Sheet 001", ADD, "Size from load list; add battery packs if > 30 min autonomy is required."),
  ("UPS-01", "AP9641", "UPS Network Management Card 3 (SNMP / Modbus TCP)", 1, "pc", "Sheet 002", ADD, "Address 10.10.30.61."),
  ("RACK-CR", "AR3100", "NetShelter SX 42U 600 × 1070 mm rack – server, SW1, FW-01, GPS-01, UPS-01", 1, "pc", "Sheet 001", ADD, ""),
  ("CONSOLE", "TBC-CONSOLE", "Operator console, 2 positions, 2 × 27\" monitors each, keyboard / mouse", 1, "lot", "Sheet 001", TBC, "May be by building contractor."),
 ]),
 ("B", "Main PLC panel – Drop 0 (control room)", [
  ("PLC-MAIN", "BMEXBP0800", "PLC 8-slot Ethernet backplane", 1, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "BMXCPS3020", "X80 DC power supply 24–48 V DC, 32 W, isolated", 1, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "BMEP582040", "M580 CPU level 2, remote and distributed I/O – plant controller (PPC)", 1, "pc", "Offer", OFF, "Second offer CPU is the spare (location S)."),
  ("PLC-MAIN", "BMENOC0311", "M580 I/O scanner module, FactoryCast, EtherNet/IP & Modbus TCP, 3 ports", 2, "pc", "Sheet 001", OFF, "#1 SCADA LAN, #2 second SCADA LAN path. Offer has 1 in use (+1)."),
  ("PLC-MAIN", "BMENOC0321", "M580 I/O scanner module, IP forwarding, coated – device network", 1, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "BMXDDI3202K", "X80 digital input module 32 ch 24 V DC sink", 1, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "BMXDDO3202K", "X80 digital output module 32 ch 24 V DC source", 1, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "BMXFCW303", "3 m pre-equipped cable, 1 connector + free wires (one per 32-ch module)", 2, "pc", "Offer", OFF, ""),
  ("PLC-MAIN", "ABL8RPS24100", "Phaseo power supply 24 V DC 10 A", 2, "pc", "Panel design", ADD, "Redundant 24 V DC for PLC and I/O."),
  ("PLC-MAIN", "ABL8RED24400", "Phaseo redundancy module 24 V DC 40 A", 1, "pc", "Panel design", ADD, ""),
  ("PLC-MAIN", "RSB1A120BD", "Zelio RSB interposing relay 1 C/O 12 A, 24 V DC coil", 8, "pc", "Panel design", ADD, "DO isolation to field circuits."),
  ("PLC-MAIN", "RSZE1S35M", "Socket for RSB1A relay, with clip", 8, "pc", "Panel design", ADD, ""),
  ("PLC-MAIN", "NSYTRV42", "Linergy TR screw terminal block 4 mm², feed-through", 150, "pc", "Panel design", ADD, "Estimate; confirm from wiring diagrams."),
  ("PLC-MAIN", "TBC-ENC-PLC", "Floor-standing enclosure 2000 × 800 × 600 IP55, fan, light, socket, earthing", 1, "pc", "Panel design", TBC, "Spacial SM or equivalent."),
 ]),
 ("C", "Rack 2 panel – Drop 1, MV / BESS / DG I/O and device network (11 kV GIS room)", [
  ("PLC-R2", "BMEXBP1200", "PLC 12-slot Ethernet backplane", 1, "pc", "Offer", OFF, ""),
  ("PLC-R2", "BMXCPS3020", "X80 DC power supply 24–48 V DC, 32 W, isolated", 1, "pc", "Offer", OFF, ""),
  ("PLC-R2", "BMECRA31210", "X80 performance Ethernet drop adapter", 1, "pc", "Offer", OFF, ""),
  ("PLC-R2", "BMXDDI3202K", "X80 digital input module 32 ch 24 V DC sink", 5, "pc", "I/O list", OFF, "Offer 3; +2 for 138 hard-wired DI + 20 % spare (see IO Capacity)."),
  ("PLC-R2", "BMXDDO3202K", "X80 digital output module 32 ch 24 V DC source", 2, "pc", "Offer", OFF, ""),
  ("PLC-R2", "BMXFCW303", "3 m pre-equipped cable, 1 connector + free wires (one per 32-ch module)", 7, "pc", "I/O list", OFF, "One per module (5 DI + 2 DO)."),
  ("PLC-R2", "BMXNOM0200", "X80 serial link module, Modbus RTU, 2 × RS485 – DG controllers", 1, "pc", "Offer", OFF, ""),
  ("PLC-R2", "VW3A8306R10", "RS485 cordset 1 m, 2 × RJ45", 1, "pc", "Offer", OFF, ""),
  ("SW2", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – device network (SW2)", 1, "pc", "Sheet 002", TBC, "Replaces offer MCSESM083F23F1 (copper only) so SW2 can close the fibre ring."),
  ("PLC-R2", "ABL8RPS24100", "Phaseo power supply 24 V DC 10 A", 2, "pc", "Panel design", ADD, ""),
  ("PLC-R2", "ABL8RED24400", "Phaseo redundancy module 24 V DC 40 A", 1, "pc", "Panel design", ADD, ""),
  ("PLC-R2", "RSB1A120BD", "Zelio RSB interposing relay 1 C/O 12 A, 24 V DC coil – VCB / ACB / DG commands", 32, "pc", "I/O list", ADD, "26 DO in I/O list + spare."),
  ("PLC-R2", "RSZE1S35M", "Socket for RSB1A relay, with clip", 32, "pc", "Panel design", ADD, ""),
  ("PLC-R2", "NSYTRV42", "Linergy TR screw terminal block 4 mm², feed-through", 400, "pc", "Panel design", ADD, "Estimate; 164 DI/DO + commons + supplies."),
  ("PLC-R2", "TBC-ENC-PLC", "Floor-standing enclosure 2000 × 800 × 600 IP55, fan, light, socket, earthing", 1, "pc", "Panel design", TBC, ""),
 ]),
 ("D", "11 kV GIS combiner panel – protection relays", [
  ("P5M30-AH1…AH8", "REL50458", "Easergy P5M30 48–250 V, 3 CT, 1 CSH, 4 VT, 16 DI / 12 DO, RSTP, Ethernet RJ45", 8, "pc", "Offer (PDF)", OFF, "Price not in offer workbook. Confirm GIS panel vendor does not already supply relays."),
 ]),
 ("E1", "President office TX station – LVDB 1 (NTS-LVP1/LVP2)", [
  ("NTS-LVP1/2", "NTSNEC1200K", "Edge I/O NTS network interface module + base + termination, EtherNet/IP", 2, "pc", "Offer", OFF, ""),
  ("NTS-LVP1/2", "NTSPFB1002HK", "Edge I/O NTS power supply module + base, 24 V DC, bus and field", 2, "pc", "Offer", OFF, ""),
  ("NTS-LVP1/2", "NTSXTB02030H", "Screw terminal block 2 pt, 5 mm, for power supply module", 4, "pc", "Offer", OFF, "2 per power supply module."),
  ("NTS-LVP1/2", "NTSDDI1602XK", "Edge I/O NTS discrete input 16 ch 24 V DC sink + base", 4, "pc", "Offer", OFF, "29 DI required."),
  ("NTS-LVP1/2", "NTSXTB18001XH", "Screw terminal block 18 pt, 5 mm, high-height (DI)", 4, "pc", "Offer", OFF, "One per DI module (offer had 3)."),
  ("NTS-LVP1/2", "NTSDDO1602K", "Edge I/O NTS discrete output 16 ch 24 V DC 0.5 A + base", 2, "pc", "Offer", OFF, "11 DO required (incomer + 10 shunt trips)."),
  ("NTS-LVP1/2", "NTSXTB18001H", "Screw terminal block 18 pt, 3.81 mm, low-height (DO)", 2, "pc", "Offer", OFF, "One per DO module (offer had 1)."),
  ("NTS-LVP1", "NTSFMB0120K", "Edge I/O NTS field device master, RS485 / RS232 Modbus RTU", 1, "pc", "Offer", OFF, "Spare serial port for future meters."),
  ("NTS-LVP1", "NTSXTB12001H", "Screw terminal block 12 pt, 5 mm (serial module)", 1, "pc", "Offer", OFF, ""),
  ("PAS600-LVP1/2", "PAS600", "EcoStruxure Universal Panel Server 110–277 V AC/DC", 2, "pc", "Offer", OFF, ""),
  ("PAS600-LVP1/2", "130729", "Asset Link humidity sensor CL110", 2, "pc", "Offer", OFF, ""),
  ("PAS600-LVP1/2", "EMS59440", "Asset Link temperature sensor TH110", 4, "pc", "Offer", OFF, ""),
  ("PM5320-LVP", "METSEPM5320", "PM5320 power meter, Ethernet, harmonics to 31st, 2 DI / 2 DO", 1, "pc", "Offer (PDF)", OFF, "Price not in offer workbook."),
  ("LS-LVP", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – field switch", 1, "pc", "Sheet 002", TBC, ""),
  ("NTS-LVP", "TBC-ENC-RIO", "Wall-mounted enclosure 800 × 600 × 300 IP55 for Edge I/O, switch and PSU", 1, "pc", "Panel design", TBC, ""),
  ("NTS-LVP", "ABL8RPS24050", "Phaseo power supply 24 V DC 5 A", 1, "pc", "Panel design", ADD, "Feed from RMU UPS 2 kVA."),
  ("NTS-LVP", "RSB1A120BD", "Zelio RSB interposing relay 1 C/O 12 A, 24 V DC coil – shunt trips", 12, "pc", "I/O list", ADD, ""),
  ("NTS-LVP", "RSZE1S35M", "Socket for RSB1A relay, with clip", 12, "pc", "Panel design", ADD, ""),
 ]),
 ("E2", "Prime Minister office TX station – LVDB 2 (NTS-LVM1)", [
  ("NTS-LVM1", "NTSNEC1200K", "Edge I/O NTS network interface module + base + termination, EtherNet/IP", 1, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSPFB1002HK", "Edge I/O NTS power supply module + base, 24 V DC, bus and field", 1, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSXTB02030H", "Screw terminal block 2 pt, 5 mm, for power supply module", 2, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSDDI1602XK", "Edge I/O NTS discrete input 16 ch 24 V DC sink + base", 3, "pc", "Offer", OFF, "29 DI required."),
  ("NTS-LVM1", "NTSXTB18001XH", "Screw terminal block 18 pt, 5 mm, high-height (DI)", 3, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSDDO1602K", "Edge I/O NTS discrete output 16 ch 24 V DC 0.5 A + base", 1, "pc", "Offer", OFF, "11 DO required."),
  ("NTS-LVM1", "NTSXTB18001H", "Screw terminal block 18 pt, 3.81 mm, low-height (DO)", 1, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSFMB0120K", "Edge I/O NTS field device master, RS485 / RS232 Modbus RTU", 1, "pc", "Offer", OFF, ""),
  ("NTS-LVM1", "NTSXTB12001H", "Screw terminal block 12 pt, 5 mm (serial module)", 1, "pc", "Offer", OFF, ""),
  ("PAS600-LVM1", "PAS600", "EcoStruxure Universal Panel Server 110–277 V AC/DC", 1, "pc", "Offer", OFF, ""),
  ("PAS600-LVM1", "130729", "Asset Link humidity sensor CL110", 1, "pc", "Offer", OFF, ""),
  ("PAS600-LVM1", "EMS59440", "Asset Link temperature sensor TH110", 1, "pc", "Offer", OFF, ""),
  ("PM5320-LVM", "METSEPM5320", "PM5320 power meter, Ethernet, harmonics to 31st, 2 DI / 2 DO", 1, "pc", "Offer (PDF)", OFF, ""),
  ("LS-LVM", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – field switch", 1, "pc", "Sheet 002", TBC, ""),
  ("NTS-LVM", "TBC-ENC-RIO", "Wall-mounted enclosure 800 × 600 × 300 IP55 for Edge I/O, switch and PSU", 1, "pc", "Panel design", TBC, ""),
  ("NTS-LVM", "ABL8RPS24050", "Phaseo power supply 24 V DC 5 A", 1, "pc", "Panel design", ADD, ""),
  ("NTS-LVM", "RSB1A120BD", "Zelio RSB interposing relay 1 C/O 12 A, 24 V DC coil – shunt trips", 12, "pc", "I/O list", ADD, ""),
  ("NTS-LVM", "RSZE1S35M", "Socket for RSB1A relay, with clip", 12, "pc", "Panel design", ADD, ""),
 ]),
 ("F", "PV inverter station – PV areas 1–4 (NTS-PV1…PV4)", [
  ("NTS-PV1…4", "NTSNEC1200K", "Edge I/O NTS network interface module + base + termination, EtherNet/IP", 4, "pc", "Offer", OFF, "One island per PV area."),
  ("NTS-PV1…4", "NTSPFB1002HK", "Edge I/O NTS power supply module + base, 24 V DC, bus and field", 4, "pc", "Offer", OFF, ""),
  ("NTS-PV1…4", "NTSXTB02030H", "Screw terminal block 2 pt, 5 mm, for power supply module", 8, "pc", "Offer", OFF, ""),
  ("NTS-PV1…4", "NTSDDI1602XK", "Edge I/O NTS discrete input 16 ch 24 V DC sink + base", 12, "pc", "Offer", OFF, "Offer 3 per area; 4 DI needed per area – see Changes, option O1."),
  ("NTS-PV1…4", "NTSXTB18001XH", "Screw terminal block 18 pt, 5 mm, high-height (DI)", 12, "pc", "Offer", OFF, ""),
  ("NTS-PV1…4", "NTSDDO1602K", "Edge I/O NTS discrete output 16 ch 24 V DC 0.5 A + base", 4, "pc", "Offer", OFF, "No hard-wired DO needed in PV areas – see option O1."),
  ("NTS-PV1…4", "NTSXTB18001H", "Screw terminal block 18 pt, 3.81 mm, low-height (DO)", 4, "pc", "Offer", OFF, ""),
  ("NTS-PV1…4", "NTSFMB0120K", "Edge I/O NTS field device master, RS485 Modbus RTU – one SG350HX each", 4, "pc", "Offer", OFF, "Point-to-point to each inverter (sheet 002, note 4)."),
  ("NTS-PV1…4", "NTSXTB12001H", "Screw terminal block 12 pt, 5 mm (serial module)", 4, "pc", "Offer", OFF, ""),
  ("PAS600-PV1…4", "PAS600", "EcoStruxure Universal Panel Server 110–277 V AC/DC", 4, "pc", "Offer", OFF, ""),
  ("PAS600-PV1…4", "130729", "Asset Link humidity sensor CL110", 4, "pc", "Offer", OFF, ""),
  ("PAS600-PV1…4", "EMS59440", "Asset Link temperature sensor TH110", 4, "pc", "Offer", OFF, ""),
  ("PM5320-PV", "METSEPM5320", "PM5320 power meter, Ethernet, harmonics to 31st, 2 DI / 2 DO", 1, "pc", "Offer (PDF)", OFF, "0.8 kV PV LV busbar; confirm CT/VT suitability at 800 V."),
  ("LS-PV", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – field switch", 1, "pc", "Sheet 002", TBC, ""),
  ("NTS-PV", "TBC-ENC-RIO", "Wall-mounted enclosure 800 × 600 × 300 IP55 for Edge I/O, switch and PSU", 2, "pc", "Panel design", TBC, "Two enclosures for 4 islands; confirm against inverter layout."),
  ("NTS-PV", "ABL8RPS24050", "Phaseo power supply 24 V DC 5 A", 2, "pc", "Panel design", ADD, ""),
 ]),
 ("G", "BESS container interface", [
  ("PM5320-BESS", "METSEPM5320", "PM5320 power meter, Ethernet, harmonics to 31st, 2 DI / 2 DO", 1, "pc", "Offer (PDF)", OFF, "0.69 kV BESS LV busbar."),
  ("LS-BESS", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – field switch", 1, "pc", "Sheet 002", TBC, "Mount in vendor-provided space in container."),
 ]),
 ("H", "Diesel generator area", [
  ("PM5320-DG", "METSEPM5320", "PM5320 power meter, Ethernet, harmonics to 31st, 2 DI / 2 DO", 1, "pc", "Offer (PDF)", OFF, "DG 0.4 kV common bus."),
  ("LS-DG", "TBC-MCSESM-FO", "Modicon managed switch 8TX + 2 FX single-mode – field switch", 1, "pc", "Sheet 002", TBC, ""),
 ]),
 ("J", "Communication cabling and infrastructure", [
  ("FO ring", "TBC-FO-OS2-4C", "Fibre optic cable single-mode OS2, 4 core, armoured, outdoor", 1500, "m", "Sheet 002", TBC, "Estimate – take lengths from LV cable route drawings."),
  ("FO ring", "TBC-FO-PP", "Fibre patch panel / splice box, 6 × LC duplex, with pigtails", 6, "pc", "Sheet 002", TBC, "Control room + 5 field switches."),
  ("FO ring", "TBC-FO-PC", "Fibre patch cord LC–LC duplex single-mode 2 m", 12, "pc", "Sheet 002", TBC, ""),
  ("Ethernet", "TBC-CAT6", "Cat6 S/FTP industrial Ethernet cable", 500, "m", "Sheet 002", TBC, "Estimate."),
  ("Ethernet", "TBC-RJ45", "RJ45 shielded field-installable connector, Cat6", 100, "pc", "Sheet 002", TBC, ""),
  ("RS485", "TBC-RS485", "RS485 cable 2-pair shielded twisted pair 120 Ω (Belden 3106A or equal)", 600, "m", "Sheet 002", TBC, "4 inverters + DG controllers; estimate."),
  ("Hard-wired I/O", "TBC-MULTICORE", "Screened multicore 1.0 mm² control cable for DI / DO", 1, "lot", "Sheet 001", TBC, "Often by electrical contractor – confirm."),
 ]),
 ("K", "Third-party interfaces (by others – listed for completeness)", [
  ("INV NB01…04", "SG350HX-RS485", "Sungrow SG350HX RS485 Modbus RTU port and register map", 4, "pc", "Sheet 003", BYO, "PV EPC scope."),
  ("BESS-EMS", "BESS-EMS-MBTCP", "BESS EMS Modbus TCP server, register map and heartbeat", 1, "lot", "Sheet 004", BYO, "BESS vendor scope."),
  ("DGC-1/2", "DGC-MBRTU", "Genset controller Modbus RTU communication module", 2, "pc", "Sheet 001", BYO, "DG vendor scope; controller make TBC."),
  ("WS-01", "BLUELOG-XM5000", "Meteocontrol blue'Log XM 5000 weather station with sensors", 1, "lot", "Sheet 003", BYO, "PV EPC scope (drawing Weather Station Mast)."),
  ("UPS-02/03", "RMU-UPS-CARD", "Communication card for RMU 2 kVA UPS (Modbus TCP / SNMP)", 2, "pc", "Sheet 002", BYO, "RMU vendor scope; confirm availability."),
  ("FACP-GW", "TBC-FACP-GW", "Fire alarm panel Modbus gateway / volt-free output card", 1, "pc", "Sheet 002", TBC, "Fire contractor to confirm; priced as SCADA item until agreed."),
 ]),
 ("L", "Engineering and services", [
  ("—", "SVC-DESIGN-COMM", "System design and commissioning (lot, per offer)", 1, "lot", "Offer", OFF, "Confirm it includes FAT, SAT, training and as-built documents."),
  ("—", "TBC-SVC-FO", "Fibre installation, splicing and OTDR testing", 1, "lot", "Sheet 002", TBC, ""),
 ]),
 ("S", "Commissioning and operational spares", [
  ("Spare", "BMEXBP0800", "PLC 8-slot Ethernet backplane", 1, "pc", "Offer", OFF, ""),
  ("Spare", "BMXCPS3020", "X80 DC power supply 24–48 V DC", 1, "pc", "Offer", OFF, ""),
  ("Spare", "BMEP582040", "M580 CPU level 2", 1, "pc", "Offer", OFF, "Cold spare. Not a hot-standby pair."),
  ("Spare", "BMENOC0311", "M580 I/O scanner module, 3 ports", 1, "pc", "Offer", OFF, ""),
  ("Spare", "BMXDDI3202K", "X80 digital input module 32 ch", 1, "pc", "Offer", OFF, ""),
  ("Spare", "BMXDDO3202K", "X80 digital output module 32 ch", 1, "pc", "Offer", OFF, ""),
  ("Spare", "NTSNEC1200K", "Edge I/O NTS network interface module", 1, "pc", "Offer", OFF, ""),
  ("Spare", "NTSPFB1002HK", "Edge I/O NTS power supply module", 1, "pc", "Offer", OFF, ""),
  ("Spare", "NTSDDI1602XK", "Edge I/O NTS discrete input 16 ch", 1, "pc", "Offer", OFF, ""),
  ("Spare", "NTSDDO1602K", "Edge I/O NTS discrete output 16 ch", 1, "pc", "Offer", OFF, ""),
  ("Spare", "PAS600", "EcoStruxure Universal Panel Server", 1, "pc", "Offer", OFF, ""),
  ("Spare", "130729", "Asset Link humidity sensor CL110", 1, "pc", "Offer", OFF, ""),
  ("Spare", "EMS59440", "Asset Link temperature sensor TH110", 1, "pc", "Offer", OFF, ""),
  ("Spare", "TBC-MCSESM-FO", "Managed switch 8TX + 2 FX single-mode", 1, "pc", "Sheet 002", TBC, "Covers SW2 and 5 field switches."),
  ("Spare", "RSB1A120BD", "Zelio RSB interposing relay", 5, "pc", "Panel design", ADD, ""),
 ]),
]

CHANGES = {  # reference -> (change type, reason)
    "BMENOC0311": ("Quantity", "Second NOC0311 for the second SCADA LAN path (sheet 001)."),
    "BMXDDI3202K": ("Quantity", "138 hard-wired DI on the X80 racks + 20 % spare needs 6 modules; offer had 4 in use."),
    "BMXFCW303": ("Quantity", "One cable per 32-channel module (9 modules)."),
    "MCSESM083F23F1": ("Replaced", "SW2 must close the fibre ring; copper-only switch replaced by an FX variant."),
    "NTSNEC1200K": ("Quantity", "7 islands as drawn (President station has 2) + 1 spare."),
    "NTSPFB1002HK": ("Quantity", "One per island (7) + 1 spare."),
    "NTSXTB02030H": ("Quantity", "Two terminal blocks per power supply module (7 modules)."),
    "NTSDDI1602XK": ("Quantity", "Location quantities add to 19 installed; spare was not on top."),
    "NTSXTB18001XH": ("Quantity", "One terminal block per DI module (offer short by 1 at LVDB 1)."),
    "NTSDDO1602K": ("Quantity", "7 installed + 1 spare."),
    "NTSXTB18001H": ("Quantity", "One terminal block per DO module (offer short by 1 at LVDB 1)."),
    "PAS600": ("Quantity", "7 installed + 1 spare (offer spare not on top)."),
    "130729": ("Quantity", "7 installed + 1 spare."),
    "EMS59440": ("Quantity", "9 installed + 1 spare."),
    "PVSCL60E": ("Removed", "Plant uses 4 × Sungrow SG350HX supplied by the PV EPC; Schneider string inverter not required."),
}

# ------------------------------------------------------------------ styles
F = "Arial"
f_base, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_in = Font(name=F, size=10, color="0000FF")
f_title = Font(name=F, size=14, bold=True, color="1F3442")
fill_head = PatternFill("solid", fgColor="1F3442")
fill_loc = PatternFill("solid", fgColor="DCE3E8")
fill_sub = PatternFill("solid", fgColor="F3F5F6")
fill_rfq = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="B9C1C8")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
PRICE = '#,##0.00;(#,##0.00);"-"'

def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = f_head, fill_head, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ------------------------------------------------------------------ BOM by Location
ws = wb.active
ws.title = "BOM by Location"
ws["A1"] = "Villa Somalia SCADA – Bill of Material by Location (Schneider Electric architecture, drawings BFE/PR/26/01-SCADA-001/002)"
ws["A1"].font = f_title
ws["A2"] = "Blue = input quantity / price. Yellow = price to be obtained (RFQ). Prices in offer currency (to be confirmed)."
ws["A2"].font = Font(name=F, size=9, italic=True)
cols = ["Item", "Loc", "Location", "Device / tag", "Reference", "Description", "Qty", "Unit", "Unit price", "Total price", "Source", "Status", "Remarks"]
header(ws, 4, cols)
widths(ws, [6, 5, 26, 16, 20, 58, 7, 6, 12, 13, 13, 16, 52])
ws.freeze_panes = "A5"
r, item, sub_rows, line_rows = 5, 0, [], []
for code, name, lines in LOCATIONS:
    ws.cell(row=r, column=2, value=code).font = f_bold
    ws.cell(row=r, column=3, value=name).font = f_bold
    for c in range(1, 14):
        ws.cell(row=r, column=c).fill = fill_loc
        ws.cell(row=r, column=c).border = border
    r += 1
    first = r
    for tag, ref, desc, qty, unit, src, status, rem in lines:
        item += 10
        price = OFFER_PRICE.get(ref)
        if status == BYO:
            price = 0
        vals = [item, code, name, tag, ref, desc, qty, unit, price, f'=IF(ISNUMBER(I{r}),G{r}*I{r},"")', src, status, rem]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border, cell.alignment = f_base, border, wrap
        ws.cell(row=r, column=7).font = f_in
        ws.cell(row=r, column=9).font = f_in
        ws.cell(row=r, column=9).number_format = PRICE
        ws.cell(row=r, column=10).number_format = PRICE
        if price is None:
            ws.cell(row=r, column=9).fill = fill_rfq
        line_rows.append(r)
        r += 1
    ws.cell(row=r, column=6, value=f"Subtotal {code} – {name}").font = f_bold
    ws.cell(row=r, column=10, value=f"=SUBTOTAL(9,J{first}:J{r - 1})").font = f_bold
    ws.cell(row=r, column=12, value=f'=COUNTIFS(B{first}:B{r - 1},"{code}",I{first}:I{r - 1},"")&" lines RFQ"').font = Font(name=F, size=9, italic=True)
    for c in range(1, 14):
        ws.cell(row=r, column=c).fill = fill_sub
        ws.cell(row=r, column=c).border = border
    ws.cell(row=r, column=10).number_format = PRICE
    sub_rows.append((code, name, r))
    r += 2
LAST = r - 2
ws.cell(row=r, column=6, value="TOTAL – priced lines only").font = f_bold
ws.cell(row=r, column=10, value=f"=SUBTOTAL(9,J5:J{LAST})").font = f_bold
ws.cell(row=r, column=10).number_format = PRICE
ws.cell(row=r + 1, column=6, value="Lines awaiting price (RFQ)").font = f_bold
ws.cell(row=r + 1, column=10, value=f'=COUNTIFS(E5:E{LAST},"<>",I5:I{LAST},"")').font = f_bold
ws.cell(row=r + 2, column=6, value="Offer total for comparison (Schneider workbook, incl. services)").font = f_base
ws.cell(row=r + 2, column=10, value=102584.4100564402).font = f_in
ws.cell(row=r + 2, column=10).number_format = PRICE
ws.cell(row=r + 2, column=10).comment = Comment("Source: offer workbook 'Bill of Material for Villa Somalia', sheet 'BOM by Reference', Total value. Excludes P5M30 and PM5320 (no price given).", "BOM")
TOTAL_ROW = r
ws.auto_filter.ref = f"A4:M{LAST}"
RNG = lambda col: f"'BOM by Location'!${col}$5:${col}${LAST}"

# ------------------------------------------------------------------ BOM by Reference
wr = wb.create_sheet("BOM by Reference")
wr["A1"] = "Bill of Material by Reference (quantities summed from 'BOM by Location')"
wr["A1"].font = f_title
header(wr, 3, ["Reference", "Description", "Total qty", "Unit", "Unit price", "Total price", "Status", "Offer qty", "Change"])
widths(wr, [20, 62, 10, 6, 12, 14, 16, 10, 12])
wr.freeze_panes = "A4"
seen = {}
for code, name, lines in LOCATIONS:
    for tag, ref, desc, qty, unit, src, status, rem in lines:
        if ref not in seen:
            seen[ref] = (desc, unit, status)
rr = 4
for ref, (desc, unit, status) in seen.items():
    vals = [ref, desc, f"=SUMIF({RNG('E')},A{rr},{RNG('G')})", unit,
            f"=IFERROR(IF(INDEX({RNG('I')},MATCH(A{rr},{RNG('E')},0))=\"\",\"\",INDEX({RNG('I')},MATCH(A{rr},{RNG('E')},0))),\"\")",
            f'=IF(ISNUMBER(E{rr}),IF(E{rr}="","",C{rr}*E{rr}),"")', status,
            OFFER_QTY.get(ref, 0), f"=C{rr}-H{rr}"]
    for c, v in enumerate(vals, 1):
        cell = wr.cell(row=rr, column=c, value=v)
        cell.font, cell.border, cell.alignment = f_base, border, wrap
    wr.cell(row=rr, column=8).font = f_in
    for c in (5, 6):
        wr.cell(row=rr, column=c).number_format = PRICE
    wr.cell(row=rr, column=9).number_format = '+0;-0;0'
    rr += 1
wr.cell(row=rr, column=2, value="TOTAL – priced lines").font = f_bold
wr.cell(row=rr, column=6, value=f"=SUM(F4:F{rr - 1})").font = f_bold
wr.cell(row=rr, column=6).number_format = PRICE
wr.cell(row=rr + 1, column=2, value=f"Check: equals 'BOM by Location' total").font = f_base
wr.cell(row=rr + 1, column=6, value=f"=IF(ABS(F{rr}-'BOM by Location'!J{TOTAL_ROW})<0.01,\"OK\",\"MISMATCH\")").font = f_bold
wr.auto_filter.ref = f"A3:I{rr - 1}"

# ------------------------------------------------------------------ Summary by Location
wsum = wb.create_sheet("Summary")
wsum["A1"] = "Cost summary by location"
wsum["A1"].font = f_title
header(wsum, 3, ["Loc", "Location", "Priced subtotal", "Lines RFQ", "Share of priced total"])
widths(wsum, [6, 70, 16, 14, 18])
sr = 4
for code, name, row in sub_rows:
    wsum.cell(row=sr, column=1, value=code)
    wsum.cell(row=sr, column=2, value=name)
    wsum.cell(row=sr, column=3, value=f"='BOM by Location'!J{row}").number_format = PRICE
    wsum.cell(row=sr, column=4, value=f'=COUNTIFS({RNG("B")},A{sr},{RNG("I")},"",{RNG("E")},"<>")')
    wsum.cell(row=sr, column=5, value=f"=IF($C${4 + len(sub_rows)}=0,0,C{sr}/$C${4 + len(sub_rows)})").number_format = "0.0%"
    for c in range(1, 6):
        wsum.cell(row=sr, column=c).font = f_base if c != 3 else Font(name=F, size=10, color="008000")
        wsum.cell(row=sr, column=c).border = border
    sr += 1
wsum.cell(row=sr, column=2, value="Total").font = f_bold
wsum.cell(row=sr, column=3, value=f"=SUM(C4:C{sr - 1})").number_format = PRICE
wsum.cell(row=sr, column=4, value=f"=SUM(D4:D{sr - 1})")
for c in (3, 4):
    wsum.cell(row=sr, column=c).font = f_bold

# ------------------------------------------------------------------ Changes vs Offer
wc = wb.create_sheet("Changes vs Offer")
wc["A1"] = "Changes against the Schneider Electric offer (Villa Somalia, version 1)"
wc["A1"].font = f_title
header(wc, 3, ["Reference", "Description", "Offer qty", "Project qty", "Delta", "Change", "Reason"])
widths(wc, [20, 52, 10, 11, 8, 12, 78])
cr = 4
for ref in list(CHANGES) + [k for k in seen if k not in OFFER_QTY]:
    desc = seen.get(ref, ("Three-phase string inverter (Schneider)",))[0]
    ctype, reason = CHANGES.get(ref, ("Added", "Required by drawings 001/002 or panel design; not in offer."))
    vals = [ref, desc, OFFER_QTY.get(ref, 0), f"=SUMIF({RNG('E')},A{cr},{RNG('G')})", f"=D{cr}-C{cr}", ctype, reason]
    for c, v in enumerate(vals, 1):
        cell = wc.cell(row=cr, column=c, value=v)
        cell.font, cell.border, cell.alignment = f_base, border, wrap
    wc.cell(row=cr, column=3).font = f_in
    wc.cell(row=cr, column=5).number_format = '+0;-0;0'
    cr += 1
cr += 1
wc.cell(row=cr, column=1, value="Options not applied (decision needed)").font = f_bold
opts = [
    ("O1", "PV areas: 4 hard-wired DI per area need 1 × NTSDDI1602XK and no DO module. Reducing from 3 DI + 1 DO per area to 1 DI saves 8 × NTSDDI1602XK, 8 × NTSXTB18001XH, 4 × NTSDDO1602K and 4 × NTSXTB18001H. Not applied because drawing 003 and the offer show 3 DI + 1 DO; keep if combiner-box signals are expected."),
    ("O2", "Hot-standby controller: replace BMEP582040 + spare by an M580 hot-standby pair (BMEH58xxxx + sync link) if the client requires controller redundancy."),
    ("O3", "Relays P5M30 and PM5320 meters: remove from this BOM if the GIS panel / LV board vendors already supply them."),
]
for k, t in opts:
    cr += 1
    wc.cell(row=cr, column=1, value=k).font = f_bold
    wc.cell(row=cr, column=2, value=t).font = f_base
    wc.merge_cells(start_row=cr, start_column=2, end_row=cr, end_column=7)
    wc.cell(row=cr, column=2).alignment = wrap
    wc.row_dimensions[cr].height = 42

# ------------------------------------------------------------------ IO Capacity
wi = wb.create_sheet("IO Capacity")
wi["A1"] = "Hard-wired I/O capacity check"
wi["A1"].font = f_title
wi["A3"], wi["B3"] = "Spare required", 0.2
wi["A3"].font, wi["B3"].font = f_bold, f_in
wi["B3"].number_format = "0%"
wi["C3"] = "Input. Required DI/DO counts come from docs/villa-somalia-io-list.csv (types DI and DO)."
wi["C3"].font = Font(name=F, size=9, italic=True)
header(wi, 5, ["Area", "Locations", "Type", "Required (I/O list)", "Required incl. spare", "Module", "Channels / module", "Modules installed", "Channels installed", "Free channels", "Free %", "Check"])
widths(wi, [34, 10, 6, 12, 12, 16, 11, 11, 11, 11, 9, 9])
areas = [
    ("X80 racks (main + Rack 2)", ["B", "C"], "DI", 138, "BMXDDI3202K", 32),
    ("X80 racks (main + Rack 2)", ["B", "C"], "DO", 26, "BMXDDO3202K", 32),
    ("President TX station", ["E1"], "DI", 29, "NTSDDI1602XK", 16),
    ("President TX station", ["E1"], "DO", 11, "NTSDDO1602K", 16),
    ("Prime Minister TX station", ["E2"], "DI", 29, "NTSDDI1602XK", 16),
    ("Prime Minister TX station", ["E2"], "DO", 11, "NTSDDO1602K", 16),
    ("PV areas 1–4 (total)", ["F"], "DI", 16, "NTSDDI1602XK", 16),
    ("PV areas 1–4 (total)", ["F"], "DO", 0, "NTSDDO1602K", 16),
]
ir = 6
for name, locs, typ, req, mod, ch in areas:
    inst = "+".join(f'SUMIFS({RNG("G")},{RNG("E")},F{ir},{RNG("B")},"{l}")' for l in locs)
    vals = [name, ", ".join(locs), typ, req, f"=ROUNDUP(D{ir}*(1+$B$3),0)", mod, ch, f"={inst}", f"=G{ir}*H{ir}",
            f"=I{ir}-D{ir}", f"=IF(I{ir}=0,0,J{ir}/I{ir})", f'=IF(I{ir}>=E{ir},"OK","SHORT")']
    for c, v in enumerate(vals, 1):
        cell = wi.cell(row=ir, column=c, value=v)
        cell.font, cell.border = f_base, border
    wi.cell(row=ir, column=4).font = f_in
    wi.cell(row=ir, column=11).number_format = "0%"
    ir += 1

# ------------------------------------------------------------------ Read Me (first)
rm = wb.create_sheet("Read Me", 0)
widths(rm, [26, 100])
rows = [
    ("Villa Somalia SCADA – Project Bill of Material", None),
    ("Project", "Solarization Villa Somalia – Power Plant Control & SCADA System"),
    ("Client", "Ministry of Energy & Water Resources – Somalia"),
    ("Contractor", "Blueflag Energy Co. Ltd"),
    ("Platform", "Schneider Electric: Modicon M580 + X80, Modicon Edge I/O NTS, Easergy P5M30, PowerLogic PM5320, PAS600, EcoStruxure Machine SCADA Expert"),
    ("Basis", "Drawings BFE/PR/26/01-SCADA-001 (architecture) and -002 (communication network), I/O list docs/villa-somalia-io-list.csv, Schneider offer 'Villa Somalia' v1 and its BOM workbook."),
    ("Revision", "Rev A – draft for internal review, 22-OCT-2026"),
    ("", ""),
    ("How to read", None),
    ("BOM by Location", "Every line by panel / location, with source, status and remarks. Subtotals per location."),
    ("BOM by Reference", "Same items summed per part number, with offer quantity and difference."),
    ("Summary", "Priced subtotal and open RFQ lines per location."),
    ("Changes vs Offer", "Every quantity change, replacement, removal and addition against the Schneider offer, plus options needing a decision."),
    ("IO Capacity", "Hard-wired DI/DO required by the I/O list against channels installed, with the spare margin as an input."),
    ("", ""),
    ("Cell colours", None),
    ("Blue text", "Input: quantity, unit price or count you can change."),
    ("Yellow fill", "Unit price not available yet – request a quotation (RFQ)."),
    ("Green text", "Value linked from another sheet."),
    ("", ""),
    ("Status values", None),
    ("Offered", "In the Schneider offer (quantity may be corrected)."),
    ("Added", "Needed by the drawings or panel design; part number is a catalogue reference."),
    ("Added – ref. TBC", "Needed; exact part number to be selected / confirmed with the supplier."),
    ("By others", "Supplied by another package vendor; listed so the interface is not missed. Price 0."),
    ("", ""),
    ("Key open points", None),
    ("Prices", "Unit prices are copied from the offer workbook. P5M30 and PM5320 had no price there. Currency to be confirmed."),
    ("Field switches", "Five field switches and SW2 need 2 single-mode fibre ports; MCSESM083F23F1 is copper only. Select the fibre variant with Schneider."),
    ("Second CPU", "The offer's second BMEP582040 is a spare. It is not a hot-standby pair (see option O2)."),
    ("PVSCL60E", "Removed: the site uses Sungrow SG350HX supplied by the PV EPC."),
    ("Cable lengths", "Fibre, Ethernet and RS485 lengths are estimates; confirm from the cable route drawings."),
]
for i, (a, b) in enumerate(rows, 1):
    rm.cell(row=i, column=1, value=a)
    rm.cell(row=i, column=2, value=b)
    if b is None:
        rm.cell(row=i, column=1).font = f_title if i == 1 else f_bold
    else:
        rm.cell(row=i, column=1).font = f_bold
        rm.cell(row=i, column=2).font = f_base
        rm.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
rm["A17"].font = Font(name=F, size=10, bold=True, color="0000FF")
rm["A18"].fill = fill_rfq
rm["A19"].font = Font(name=F, size=10, bold=True, color="008000")

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = sheet.title in ("BOM by Location", "BOM by Reference")
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("wrote", OUT)
